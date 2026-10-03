"""Documents attached to a thread (bambooai/documents.py): parsing to units with locators for each
format, the map, text.md, the thread folder and its manifest, ids never reused, the refusals, the
kernel sync and the prompt block. Fixtures are generated here, not checked in.
Run: python3 tests/analyst/test_documents.py"""
import os, sys, json, tempfile, shutil
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path[:0] = [os.path.dirname(os.path.abspath(__file__)), ROOT]
import _stubs  # noqa: F401  (the hosted-only packages faked where absent, as every suite does)
from bambooai import documents as D

passed, failed = [], []
def check(name, cond, detail=""):
    (passed if cond else failed).append(name); print(("  ok   " if cond else "  FAIL ") + name + ("" if cond else f"  <- {str(detail)[:400]}"))

tmp = tempfile.mkdtemp(prefix="bamboo_docs_")

# ---- fixtures ----
def make_pdf(path):
    """A three-page paper through matplotlib's PdfPages (a dependency everywhere; TrueType fonts so the
    text is extractable): a title and headings by size, body text, a caption and a ruled table, a figure caption."""
    import logging, textwrap
    logging.getLogger("fontTools").setLevel(logging.WARNING)
    import matplotlib
    matplotlib.use("Agg"); matplotlib.rcParams["pdf.fonttype"] = 42
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages
    lorem = ("Masters athletes who kept a steady weekly load reported shorter recovery after hard sessions. "
             "The association held after adjusting for age and sleep, and it was weaker for the oldest group.")
    def page(pdf, blocks, table=None, table_y=None):
        fig = plt.figure(figsize=(8.27, 11.69)); y = 0.93
        for size, text, gap in blocks:
            for ln in textwrap.wrap(text, 95 if size <= 10 else 60):
                fig.text(0.1, y, ln, fontsize=size, family="DejaVu Sans"); y -= (size * 1.3) / (11.69 * 72)
            y -= gap / (11.69 * 72)
        if table:
            ax = fig.add_axes([0.1, table_y, 0.6, 0.12]); ax.axis("off")
            ax.table(cellText=table[1:], colLabels=table[0], loc="center", cellLoc="center")
        pdf.savefig(fig); plt.close(fig)
    with PdfPages(path) as pdf:
        page(pdf, [(18, "Training Load and Recovery in Masters Athletes", 10), (16, "Abstract", 4), (10, lorem, 8),
                   (10, lorem + " " + lorem, 0)])
        page(pdf, [(16, "2 Methods", 4), (13, "2.1 Participants", 4),
                   (10, "Forty-two athletes aged 40 to 68 were followed for twelve weeks. " + lorem, 8),
                   (10, "Table 1. Participant characteristics", 0)],
             table=[["Group", "n", "Age", "Weekly hours"], ["40-49", "15", "44.1", "7.2"], ["50-59", "16", "54.3", "6.8"],
                    ["60-68", "11", "63.0", "5.9"]], table_y=0.62)
        page(pdf, [(16, "3 Results", 4), (10, lorem, 8), (10, "Figure 1. Weekly load by athlete and age group", 0)])

def make_blank_pdf(path):
    """Two pages with a drawn rectangle and no text: the shape of a scan without a text layer."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import Rectangle
    from matplotlib.backends.backend_pdf import PdfPages
    with PdfPages(path) as pdf:
        for _ in range(2):
            fig = plt.figure(figsize=(8.27, 11.69)); ax = fig.add_axes([0, 0, 1, 1]); ax.axis("off")
            ax.add_patch(Rectangle((0.15, 0.1), 0.6, 0.7, fill=False)); pdf.savefig(fig); plt.close(fig)

def make_docx(path):
    import docx
    d = docx.Document()
    d.add_heading("Sales meeting 30 September", 0)
    d.add_heading("Agenda", 1)
    d.add_paragraph("The meeting covered the Q3 pipeline, the APAC renewals and the hiring plan.")
    d.add_heading("Q3 pipeline", 2)
    d.add_paragraph("Maria said the Q3 pipeline is 4.2 million, up from 3.1 million at the end of Q2, "
                    "with two deals above half a million each.")
    d.add_paragraph("EMEA renewals", style="List Bullet")
    d.add_paragraph("APAC expansion", style="List Bullet")
    d.add_paragraph("Table 1. Pipeline by region", style="Caption")
    t = d.add_table(rows=3, cols=3)
    for r, row in enumerate([["Region", "Pipeline", "Deals"], ["EMEA", "2.1M", "14"], ["APAC", "2.1M", "9"]]):
        for c, val in enumerate(row):
            t.cell(r, c).text = val
    d.add_heading("Hiring", 1)
    d.add_paragraph("Two account executives to be hired in Q4.")
    d.save(path)

MD = """# Post-session interview - 2026-09-28

## Context
The athlete completed a 3 h ride with two 20 min threshold blocks.

## How it felt
Legs were heavy for the first hour, then fine. Sleep the night before was short, about five hours.

| Block | Power | RPE |
|---|---|---|
| 1 | 268 W | 7 |
| 2 | 255 W | 8 |

- Hydration: 1.5 L
- Food: two bars

```
raw note: HR drifted 6% in block 2
```

Figure 1. Power and heart rate over the ride
"""

TXT = "First paragraph of plain notes about the dataset.\n\nSecond paragraph, with a number: 42 rows were dropped.\n\nThird paragraph.\n"

pdf_path, blank_pdf, docx_path, md_path, txt_path = [os.path.join(tmp, n) for n in ("paper.pdf", "scan.pdf", "meeting.docx", "interview.md", "notes.txt")]
make_pdf(pdf_path); make_blank_pdf(blank_pdf); make_docx(docx_path)
open(md_path, "w", encoding="utf-8").write(MD); open(txt_path, "w", encoding="utf-8").write(TXT)

# ---- the checks that need no parsing ----
def refused(fn, *a, **k):
    try:
        fn(*a, **k); return None
    except D.Refused as exc:
        return str(exc)
check("check_upload: a type outside pdf/docx/md/txt is refused by name", (refused(D.check_upload, "notes.rtf", 10, 0) or "").startswith("notes.rtf is not a PDF"), refused(D.check_upload, "notes.rtf", 10, 0))
check("check_upload: over 10 MB is refused with both sizes stated", refused(D.check_upload, "big.pdf", 14 * 1024 * 1024, 0) == "big.pdf is 14.0 MB; the limit is 10 MB a document.", refused(D.check_upload, "big.pdf", 14 * 1024 * 1024, 0))
check("check_upload: a fifth document is refused", refused(D.check_upload, "e.pdf", 10, 4) == "Maximum 4 documents per thread.")

# ---- PDF ----
p = D.parse(pdf_path, "paper.pdf")
kinds = [u["kind"] for u in p["units"]]
heads = [u["text"] for u in p["units"] if u["kind"] == "heading"]
tables = [u for u in p["units"] if u["kind"] == "table"]
caps = [u["text"] for u in p["units"] if u["kind"] == "caption"]
check("pdf: three pages, units numbered from 1", p["pages"] == 3 and p["units"][0]["n"] == 1 and p["units"][-1]["n"] == len(p["units"]), (p["pages"], len(p["units"])))
check("pdf: headings found by size (title, Abstract, 2 Methods, 2.1 Participants, 3 Results)",
      all(any(h.startswith(x) for h in heads) for x in ("Training Load", "Abstract", "2 Methods", "2.1 Participants", "3 Results")), heads)
check("pdf: paragraphs carry their page", any(u["kind"] == "paragraph" and u["page"] == 1 and "Masters athletes" in u["text"] for u in p["units"]), [(u["kind"], u["page"]) for u in p["units"]][:8])
check("pdf: the table is a unit with its rows, on page 2, numbered table 1", len(tables) == 1 and tables[0]["page"] == 2 and tables[0]["table_no"] == 1 and tables[0]["rows"][0][0] == "Group" and len(tables[0]["rows"]) == 4, tables and tables[0]["rows"])
check("pdf: the caption 'Table 1. ...' beside the table names it and stays a unit", tables and tables[0].get("caption", "").startswith("Table 1.") and any(c.startswith("Table 1.") for c in caps), (tables and tables[0].get("caption"), caps))
check("pdf: the figure caption is a caption unit on page 3", any(c.startswith("Figure 1.") for c in caps) and any(u["kind"] == "caption" and u["page"] == 3 for u in p["units"]), caps)
check("pdf: the table's text is not repeated as paragraphs", not any(u["kind"] == "paragraph" and "40-49" in u["text"] for u in p["units"]), [u["text"][:40] for u in p["units"] if u["page"] == 2])
check("pdf: a blank PDF is refused for having no text layer", "no text layer" in (refused(D.parse, blank_pdf, "scan.pdf") or ""), refused(D.parse, blank_pdf, "scan.pdf"))

# ---- Word ----
w = D.parse(docx_path, "meeting.docx")
wh = [(u["text"], u.get("level"), u["section"]) for u in w["units"] if u["kind"] == "heading"]
check("docx: headings with levels and the trail as section", ("Agenda", 1, "Sales meeting 30 September > Agenda") in wh and ("Q3 pipeline", 2, "Sales meeting 30 September > Agenda > Q3 pipeline") in wh, wh)
check("docx: a paragraph's section is the trail above it", any(u["kind"] == "paragraph" and "4.2 million" in u["text"] and u["section"].endswith("Q3 pipeline") for u in w["units"]), [(u["kind"], u["section"]) for u in w["units"]])
check("docx: consecutive list paragraphs become one list unit", any(u["kind"] == "list" and "- EMEA renewals\n- APAC expansion" == u["text"] for u in w["units"]), [u["text"] for u in w["units"] if u["kind"] == "list"])
wt = [u for u in w["units"] if u["kind"] == "table"]
check("docx: the table with its Caption-styled caption", len(wt) == 1 and wt[0]["rows"][1] == ["EMEA", "2.1M", "14"] and wt[0].get("caption", "").startswith("Table 1."), wt)
check("docx: a heading after the table closes the trail (Hiring at level 1)", any(u["kind"] == "paragraph" and "account executives" in u["text"] and u["section"] == "Sales meeting 30 September > Hiring" for u in w["units"]), [(u["text"][:20], u["section"]) for u in w["units"]][-3:])

# ---- Markdown and text ----
m = D.parse(md_path, "interview.md")
check("md: headings, the trail, a pipe table, a list, a fenced block, a caption",
      [u["text"] for u in m["units"] if u["kind"] == "heading"] == ["Post-session interview - 2026-09-28", "Context", "How it felt"]
      and any(u["kind"] == "table" and u["rows"][1] == ["1", "268 W", "7"] for u in m["units"])
      and any(u["kind"] == "list" and u["text"].startswith("- Hydration") for u in m["units"])
      and any(u.get("fenced") and "HR drifted" in u["text"] for u in m["units"])
      and any(u["kind"] == "caption" and u["text"].startswith("Figure 1.") for u in m["units"])
      and any(u["kind"] == "paragraph" and u["section"] == "Post-session interview - 2026-09-28 > How it felt" for u in m["units"]),
      [(u["kind"], u["text"][:30], u["section"]) for u in m["units"]])
t = D.parse(txt_path, "notes.txt")
check("txt: paragraphs by blank lines, no headings, type text", t["type"] == "text" and [u["kind"] for u in t["units"]] == ["paragraph"] * 3 and t["words"] > 10, t["units"])

# ---- locators, text.md, the map ----
check("locator: [D1.n | p.2 | table 1] for the PDF table, [D1.n | §trail | heading] for a Word heading",
      D.locator("D1", tables[0]) == f"[D1.{tables[0]['n']} | p.2 | table 1]" and D.locator("D2", [u for u in w["units"] if u["kind"] == "heading"][1]) == f"[D2.{wh and 2} | §Sales meeting 30 September > Agenda | heading]",
      (D.locator("D1", tables[0]), D.locator("D2", [u for u in w["units"] if u["kind"] == "heading"][1])))
md_text = D.render_text_md("D1", "paper.pdf", p)
check("text.md: the title line, every unit under its marker, the table as a markdown table with its csv named",
      md_text.startswith("# D1 - paper.pdf\n") and md_text.count("[D1.") == len(p["units"]) and "| Group | n | Age | Weekly hours |" in md_text and "(tables/1.csv)" in md_text, md_text[:600])
mp = D.render_map("D1", "paper.pdf", p)
check("map: one head line with type, pages, words, tables and figure captions; outline; tables with their unit; figures",
      mp.startswith("D1 - paper.pdf (PDF, 3 pages, ") and "1 table, 1 figure caption)" in mp and "Outline: Training Load" in mp and "2.1 Participants (p.2)" in mp
      and f'Tables: table 1 "Table 1. Participant characteristics" 4x4 (p.2, D1.{tables[0]["n"]})' in mp and "Figures: Figure 1. Weekly load" in mp, mp)
check("map: a few hundred tokens at most for a document of this size", len(mp) < 1200, len(mp))

# ---- the thread folder ----
tdir = os.path.join(tmp, "storage", "local", "documents", "1700000000")
e1 = D.attach(tdir, pdf_path, "paper.pdf")
e2 = D.attach(tdir, md_path, "interview.md")
man = D.load_manifest(tdir)
check("attach: D1 then D2, the manifest lists both with counts and a sha256", [d["id"] for d in man["documents"]] == ["D1", "D2"] and e1["pages"] == 3 and e1["tables"] == 1 and len(e1["sha256"]) == 64 and e2["type"] == "Markdown", man)
files = sorted(os.listdir(os.path.join(tdir, "D1"))) + sorted(os.listdir(os.path.join(tdir, "D1", "tables")))
check("attach: original, text.json, text.md, map.md and tables/1.csv are written", files == ["map.md", "original.pdf", "tables", "text.json", "text.md", "1.csv"], files)
units_json = D.read_units(tdir, "D1")
check("text.json: units carry their id D1.n and the document's name", units_json["units"][0]["id"] == "D1.1" and units_json["file"] == "paper.pdf" and units_json["units"][-1]["id"] == f"D1.{len(units_json['units'])}", units_json["units"][0])
check("attach: a refused file writes nothing and leaves the manifest as it was",
      "no text layer" in (refused(D.attach, tdir, blank_pdf, "scan.pdf") or "") and not os.path.exists(os.path.join(tdir, "D3")) and len(D.load_manifest(tdir)["documents"]) == 2, os.listdir(tdir))
D.remove(tdir, "D1")
e3 = D.attach(tdir, docx_path, "meeting.docx")
check("remove then attach: D1's folder is gone and the next id is D3, never D1 again", not os.path.exists(os.path.join(tdir, "D1")) and e3["id"] == "D3" and [d["id"] for d in D.load_manifest(tdir)["documents"]] == ["D2", "D3"], D.load_manifest(tdir))
D.attach(tdir, txt_path, "notes.txt"); D.attach(tdir, pdf_path, "paper.pdf")
check("attach: the fifth document is refused by count", refused(D.attach, tdir, txt_path, "again.txt") == "Maximum 4 documents per thread." and len(D.load_manifest(tdir)["documents"]) == 4)
big = os.path.join(tmp, "big.md"); open(big, "w").write("# Big\n\n" + ("A line of text that repeats. " * 40 + "\n\n") * 2300)
D.remove(tdir, "D4")
check("attach: parsed text over 2 MB is refused with the size stated", (refused(D.attach, tdir, big, "big.md") or "").startswith("the text of big.md is 2.") and "the limit is 2 MB" in (refused(D.attach, tdir, big, "big.md") or ""), refused(D.attach, tdir, big, "big.md"))

# ---- the kernel's copy ----
kf = D.kernel_files(tdir, D.load_manifest(tdir)["documents"][0])      # D2, the markdown with one table
check("kernel_files: text.json, text.md and the tables, never the original", [r for r, _, _ in kf] == ["D2/text.json", "D2/text.md", "D2/tables/1.csv"], kf)
send, delete = D.sync_plan([("D2/text.md", "/x", "aa"), ("D2/text.json", "/y", "bb")], {"D2/text.md": "aa", "D2/text.json": "old", "D1/text.md": "zz"})
check("sync_plan: send what is missing or changed, delete what the thread no longer has", [s[0] for s in send] == ["D2/text.json"] and delete == ["D1/text.md"], (send, delete))
kroot = os.path.join(tmp, "datasets", "local", "documents")
n_sent, n_del = D.sync_local(tdir, kroot)
inv = D.inventory(kroot)
check("sync_local: the kernel folder mirrors the manifest (D2, D3, D5), by content", n_sent == len(inv) and set(d.split("/")[0] for d in inv) == {"D2", "D3", "D5"} and "D2/tables/1.csv" in inv, inv)
D.remove(tdir, "D3"); n_sent2, n_del2 = D.sync_local(tdir, kroot)
check("sync_local: a second sync sends nothing new and removes the dropped document's files and folder", n_sent2 == 0 and n_del2 >= 2 and not os.path.exists(os.path.join(kroot, "D3")), (n_sent2, n_del2, os.listdir(kroot)))

# ---- the prompt ----
blk = D.prompt_block(tdir, "datasets/local/documents")
check("prompt_block: the where-to-find line, then every attached document's map", blk.startswith("Documents attached to this thread (2).") and "datasets/local/documents/<id>/" in blk and "D2 - interview.md (Markdown," in blk and "D5 - paper.pdf (PDF, 3 pages," in blk and "D3 -" not in blk, blk)
check("prompt_block: empty when the thread has no documents", D.prompt_block(os.path.join(tmp, "nothing"), "x") == "")
D.remove_thread(tdir)
check("remove_thread: the folder is gone", not os.path.exists(tdir))

# ---- a LaTeX-like PDF: words positioned with TJ offsets, no space glyphs (2026-10-03, an arXiv paper came out glued) ----
def tj_pdf(path, size=10, gap=-250):
    words = "Let O(t) denote the state of some subjective observer O at time t according to our lazy brain theory".split()
    arr = " ".join("(" + w.replace("(", "\\(").replace(")", "\\)") + ") " + str(gap) for w in words)
    content = ("BT /F1 %d Tf 72 700 Td [%s] TJ ET" % (size, arr)).encode()
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
            b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Times-Roman >>"]
    out = bytearray(b"%PDF-1.4\n"); offsets = []
    for i, o in enumerate(objs, 1):
        offsets.append(len(out)); out += ("%d 0 obj\n" % i).encode() + o + b"\nendobj\n"
    xref = len(out)
    out += ("xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)).encode()
    for off in offsets:
        out += ("%010d 00000 n \n" % off).encode()
    out += ("trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)).encode()
    open(path, "wb").write(out)
tj = os.path.join(tmp, "tight.pdf"); tj_pdf(tj)
tdir_tj = os.path.join(tmp, "storage", "u", "documents", "tj")
D.attach(tdir_tj, tj, "tight.pdf")
tj_text = " ".join(u.get("text", "") for u in D.read_units(tdir_tj, "D1")["units"])
check("PDF: words positioned by offsets with a 2.5pt gap and no space glyphs come out as words, not one glued string",
      "denote the state of some subjective observer" in tj_text and len(tj_text.split()) >= 18, tj_text[:120])

# ---- long paragraphs are split at sentence ends into units of at most MAX_UNIT_CHARS (2026-10-03) ----
sent = "The wet season came early and the plots on the upper terrace were sown a week after the lower ones. "
long_para = (sent * 40).strip()                              # ~4,000 characters, 40 sentences
pieces = D.split_long({"kind": "paragraph", "text": long_para, "page": 3})
check("split_long: a 4,000-character paragraph becomes three units of at most 1,500, each on its page, text preserved",
      len(pieces) == 3 and all(len(x["text"]) <= D.MAX_UNIT_CHARS for x in pieces) and all(x["page"] == 3 for x in pieces)
      and " ".join(x["text"] for x in pieces) == long_para and all(x["text"].endswith(".") for x in pieces), [len(x["text"]) for x in pieces])
check("split_long: a short paragraph, a heading, a table and a fenced block are left alone",
      D.split_long({"kind": "paragraph", "text": "short"}) == [{"kind": "paragraph", "text": "short"}] and len(D.split_long({"kind": "heading", "text": "x " * 1000})) == 1
      and len(D.split_long({"kind": "table", "rows": [["a"]], "text": ""})) == 1 and len(D.split_long({"kind": "paragraph", "text": "y " * 1000, "fenced": True})) == 1)
monster = "word " * 700                                       # one 3,500-character "sentence" with no end
mp = D.split_long({"kind": "paragraph", "text": monster.strip()})
check("split_long: a sentence longer than the cap is cut at a space", len(mp) == 3 and all(len(x["text"]) <= D.MAX_UNIT_CHARS for x in mp) and all(not x["text"].startswith(" ") for x in mp), [len(x["text"]) for x in mp])
md_long = os.path.join(tmp, "long.md"); open(md_long, "w").write("# T\n\n" + long_para + "\n")
tdir_long = os.path.join(tmp, "storage", "u", "documents", "long")
D.attach(tdir_long, md_long, "long.md")
lu = D.read_units(tdir_long, "D1")["units"]
check("parse: the split applies at upload - the units are numbered after it, ids consecutive", [u["id"] for u in lu] == ["D1.1", "D1.2", "D1.3", "D1.4"] and all(len(u.get("text", "")) <= D.MAX_UNIT_CHARS for u in lu), [u["id"] for u in lu])

# ---- the kernel's document objects (D59) ----
src_api = D.kernel_api_source(tdir_tj, "kroot")
check("kernel_api_source: defines the class, binds D1 to its folder under the kernel root and registers it in docs", "class _BambooDocument" in src_api and "D1 = _BambooDocument('D1', 'kroot/D1', 'tight.pdf'); docs['D1'] = D1" in src_api)
check("kernel_api_source: empty for a thread without documents", D.kernel_api_source(os.path.join(tmp, "nothing"), "kroot") == "")
md_api = os.path.join(tmp, "api.md"); open(md_api, "w").write("# Field notes\n\n## Weather\nThe wet season came early: rainfall was 40% above the ten-year average in March.\n\n## Regimes\nRegime B plots were sown a week later than regime A on the same soils.\n\n| plot | yield |\n|---|---|\n| 1 | 4.2 |\n| 2 | 5.1 |\n")
tdir_api = os.path.join(tmp, "storage", "u", "documents", "api"); D.attach(tdir_api, md_api, "api.md")
ns_api = {}; exec(D.kernel_api_source(tdir_api, tdir_api), ns_api); D1o = ns_api["D1"]
check("D1 in the kernel: repr names the file and the unit count; docs holds it", repr(D1o) == "D1 - api.md (6 units)" and ns_api["docs"]["D1"] is D1o, repr(D1o))
g = D1o.grep("wet season", context=1)
check("D1.grep: the matching unit with its locator and its neighbours", "[D1.3 | Field notes > Weather] The wet season came early" in g and "[D1.2 | Field notes > Weather | heading] Weather" in g and "[D1.4" in g, g)
check("D1.units / unit / around / outline", D1o.units(3, 4).startswith("[D1.3 | Field notes > Weather] The wet season") and "[D1.4" in D1o.units(3, 4) and D1o.unit("D1.3") == D1o.unit(3)
      and D1o.around("D1.4", 1).count("[D1.") == 3 and D1o.outline().splitlines()[0] == "[D1.1] Field notes" and D1o.unit(99) == "(no unit D1.99)", (D1o.outline(), D1o.around("D1.4", 1)))
check("D1.table: a DataFrame from the table's CSV; a missing table raises with the list", list(D1o.table(1).columns) == ["plot", "yield"] and len(D1o.table(1)) == 2 and D1o.text.startswith("# D1 - api.md"))
try:
    D1o.table(7); tbl_err = ""
except FileNotFoundError as exc:
    tbl_err = str(exc)
check("D1.table on a missing number says which tables exist", "D1 has no table 7" in tbl_err and "1.csv" in tbl_err, tbl_err)
check("D1.page on a document without pages says so; grep with no match says so", D1o.page(2) == "(no page 2 in D1)" and D1o.grep("zebra") == "(no match for 'zebra' in D1)")

shutil.rmtree(tmp, ignore_errors=True)
print(f"\n{len(passed)} passed, {len(failed)} failed")
sys.exit(1 if failed else 0)
